import re
import copy
import datetime
import subprocess
import tempfile
import os
import zipfile
from io import BytesIO
from pathlib import Path

from django.http import HttpResponse, FileResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAdminUser, IsAuthenticated
from rest_framework.response import Response
from docxtpl import DocxTemplate
import openpyxl

from courses.models import Course, Enrollment, Instructor
from .models import Topic, TopicFile, TopicFileProgress, QuizProgress, CourseFile, TopicQuestion, TopicAnswerChoice, TopicQuizAttempt

TEMPLATES_DIR = Path(__file__).parent / 'templates'

ALLOWED_TEMPLATES = {'oswiadczenie', 'sprzet', 'sale', 'wniosek', 'prosba', 'instruktorzy', 'prosba-recertyfikacja', 'informacja-kpp', 'sprawozdanie-egzamin-rec'}

ALLOWED_ENROLLMENT_ZIP_TEMPLATES = {'zaliczenia_tematow_KPP'}

ALLOWED_CERT_TEMPLATES = {'certyfikat'}

ALLOWED_XLSX_TEMPLATES = {'program', 'obsluga-egzaminu-rec', 'zestawienie-rec'}

ALLOWED_ATTENDANCE_XLSX_TEMPLATES = {'obecnosc'}


def _instructor_dict(inst):
    name_only = f'{inst.first_name} {inst.last_name}'.strip()
    return {
        'full_name':   str(inst),
        'name_only':   name_only,
        'title':       inst.title or '',
        'profession':  inst.profession or '',
        'specs':       inst.specializations_str,
        'years':       inst.years_experience or '',
    }


_EMPTY_INST = {'full_name': '', 'name_only': '', 'title': '', 'profession': '', 'specs': '', 'years': ''}


_STREET_PREFIX = re.compile(r'\s+(?:ul\.|al\.|pl\.|os\.|sk\.|rynek)\s+', re.IGNORECASE)

def _extract_city(exam_location):
    if not exam_location:
        return ''
    m = re.search(r'\d{2}-\d{3}\s+([^,]+)', exam_location)
    if m:
        return _STREET_PREFIX.split(m.group(1))[0].strip()
    return _STREET_PREFIX.split(exam_location.split(',')[0])[0].strip()


def _build_context(course):
    enrolled_count = course.enrollments.filter(is_deleted=False).count()

    def fmt(date):
        return date.strftime('%d.%m.%Y') if date else ''

    def fmt_str(iso):
        if not iso:
            return ''
        try:
            y, m, d = iso.split('-')
            return f'{d}.{m}.{y}'
        except Exception:
            return iso

    order = course.instructor_order or []
    if order:
        inst_map = {i.pk: i for i in Instructor.objects.filter(pk__in=order)}
        instructors_list = [_instructor_dict(inst_map[pk]) for pk in order if pk in inst_map]
    else:
        instructors_list = [_instructor_dict(i) for i in course.instructors.all()]

    # Płaskie zmienne instructor_N / instructor_N_name / _title / _profession dla slotów 1–6
    flat_instructors = {}
    for n in range(1, 7):
        d = instructors_list[n - 1] if n <= len(instructors_list) else _EMPTY_INST
        flat_instructors[f'instructor_{n}']            = d['full_name']
        flat_instructors[f'instructor_{n}_name']       = d['name_only']
        flat_instructors[f'instructor_{n}_title']      = d['title']
        flat_instructors[f'instructor_{n}_profession'] = d['profession']
        flat_instructors[f'instructor_{n}_specs']      = d['specs']
        flat_instructors[f'instructor_{n}_years']      = d['years']

    def fmt_range(start, end):
        if not start and not end:
            return ''
        if not start:
            return fmt(end)
        if not end:
            return fmt(start)
        if start.month == end.month and start.year == end.year:
            return f'{start.day} – {end.strftime("%d.%m.%Y")}'
        if start.year == end.year:
            return f'{start.strftime("%d.%m")} – {end.strftime("%d.%m.%Y")}'
        return f'{fmt(start)} – {fmt(end)}'

    return {
        'created_at':        course.created_at.strftime('%d.%m.%Y'),
        'c_n':               course.course_number or '',
        'name':              course.name,
        'course_type':       course.course_type,
        'city':              course.city,
        'address':           course.city,
        'price':             course.price,
        'max_participants':  course.max_participants,
        'enrolled_count':    enrolled_count,
        'spots_left':        course.spots_left,
        'start_date':        fmt(course.start_date),
        'end_date':          fmt(course.end_date),
        'date_range':        fmt_range(course.start_date, course.end_date),
        'exam_date':         fmt(course.exam_date),
        'exam_time':         course.exam_time.strftime('%H:%M') if course.exam_time else '',
        'exam_location':     course.exam_location or '',
        'exam_c':            _extract_city(course.exam_location),
        'y':                 str(datetime.date.today().year)[-2:],
        'course_days':       [fmt_str(d) for d in (course.course_days or [])],
        'entity_director':   course.entity_director or '',
        'academic_director': course.academic_director or '',
        # Lista słowników – w docx: {{instructors.0.full_name}}, pętla {% for i in instructors %}
        'instructors':       instructors_list,
        'psychologist':      course.psychologist or '',
        'committee_chair':   course.committee_chair or '',
        'committee_member1': course.committee_member1 or '',
        'committee_member2': course.committee_member2 or '',
        **flat_instructors,
    }


def _xlsx_replace(ws, context):
    """Podmienia {{klucz}} i {{course_days.N}} / {{instructors.N.pole}} w każdej komórce."""
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or '{{' not in cell.value:
                continue
            val = cell.value
            for key, replacement in context.items():
                if key == 'course_days':
                    for i, day in enumerate(replacement):
                        val = val.replace(f'{{{{{key}.{i}}}}}', str(day))
                elif key == 'instructors':
                    # {{instructors.N}} → full_name; {{instructors.N.pole}} → pole
                    for i, inst in enumerate(replacement):
                        val = val.replace(f'{{{{{key}.{i}}}}}', inst.get('full_name', ''))
                        for field in ('full_name', 'name_only', 'title', 'profession'):
                            val = val.replace(f'{{{{{key}.{i}.{field}}}}}', inst.get(field, ''))
                elif not isinstance(replacement, (list, dict)):
                    val = re.sub(r'\{\{\s*' + re.escape(key) + r'\s*\}\}', str(replacement), val)
            cell.value = val


def _xlsx_fill_enrollment_rows(ws, enrollments, course, ctx):
    """
    Szuka wiersza z {{p_lp}}, kopiuje jego styl dla każdego uczestnika,
    podmienia zmienne. Pozostałe wiersze – standardowy _xlsx_replace z ctx.
    """
    tpl_row_idx = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and '{{p_lp}}' in cell.value:
                tpl_row_idx = cell.row
                break
        if tpl_row_idx is not None:
            break

    # Zastąp zmienne kursu we wszystkich wierszach poza wierszem-szablonem
    for row in ws.iter_rows():
        if row[0].row == tpl_row_idx:
            continue
        for cell in row:
            if not isinstance(cell.value, str) or '{{' not in cell.value:
                continue
            val = cell.value
            for key, replacement in ctx.items():
                if not isinstance(replacement, (list, dict)):
                    val = re.sub(r'\{\{\s*' + re.escape(key) + r'\s*\}\}', str(replacement), val)
            cell.value = val

    if tpl_row_idx is None or not enrollments:
        return

    # Zapamiętaj styl wiersza-szablonu
    max_col = ws.max_column
    tpl_snapshot = []
    for col in range(1, max_col + 1):
        c = ws.cell(tpl_row_idx, col)
        tpl_snapshot.append({
            'value':         c.value,
            'font':          copy.copy(c.font),
            'fill':          copy.copy(c.fill),
            'border':        copy.copy(c.border),
            'alignment':     copy.copy(c.alignment),
            'number_format': c.number_format,
        })

    c_n   = course.course_number or ''
    year  = str(datetime.date.today().year)[-2:]
    exam_date = course.exam_date.strftime('%d.%m.%Y') if course.exam_date else ''

    for lp, enr in enumerate(enrollments, 1):
        row_idx = tpl_row_idx + lp - 1
        row_ctx = {
            'p_lp':               str(lp),
            'p_first_name':       enr.first_name or '',
            'p_last_name':        enr.last_name or '',
            'p_cert_number':      enr.cert_number or '',
            'p_cert_date':        enr.cert_date.strftime('%d.%m.%Y') if enr.cert_date else '',
            'p_new_cert_number':  f'KPP/ER/{c_n}/{year}/{str(lp).zfill(2)}',
            'p_new_cert_date':    exam_date,
        }
        for col_idx, snap in enumerate(tpl_snapshot, 1):
            cell = ws.cell(row_idx, col_idx)
            if isinstance(cell, openpyxl.cell.cell.MergedCell):
                continue
            cell.font          = copy.copy(snap['font'])
            cell.fill          = copy.copy(snap['fill'])
            cell.border        = copy.copy(snap['border'])
            cell.alignment     = copy.copy(snap['alignment'])
            cell.number_format = snap['number_format']
            val = snap['value']
            if isinstance(val, str):
                for key, replacement in row_ctx.items():
                    val = val.replace(f'{{{{{key}}}}}', str(replacement))
            cell.value = val


def _xlsx_fill_obsluga_egzaminu_rec(wb, enrollments, course, ctx):
    """Wypełnia arkusz DANE danymi kursu i uczestników; Arkusz1 – szablon wierszy."""
    ws_dane = wb['DANE ']

    # Daty zjazdów w N2–N7, data egzaminu w N8 (kolumna N = 14)
    days = ctx.get('course_days', [])
    for i, day in enumerate(days[:6]):
        ws_dane.cell(row=2 + i, column=14).value = day
    ws_dane.cell(row=8, column=14).value = ctx.get('exam_date', '')

    # Komisja egzaminacyjna (M = 13)
    ws_dane.cell(row=10, column=13).value = ctx.get('committee_chair', '')
    ws_dane.cell(row=11, column=13).value = ctx.get('committee_member1', '')
    ws_dane.cell(row=12, column=13).value = ctx.get('committee_member2', '')

    # Czyść stare dane uczestników (wiersze 3–29)
    for r in range(3, 30):
        for col in (1, 2, 3, 4, 8, 9, 10):
            ws_dane.cell(row=r, column=col).value = None

    # Wpisz uczestników
    for i, enr in enumerate(enrollments, 1):
        r = 2 + i
        ws_dane.cell(row=r, column=1).value = str(i).zfill(2)
        ws_dane.cell(row=r, column=2).value = enr.last_name or ''
        ws_dane.cell(row=r, column=3).value = enr.first_name or ''
        ws_dane.cell(row=r, column=4).value = enr.phone or ''
        ws_dane.cell(row=r, column=8).value = enr.pesel or ''
        ws_dane.cell(row=r, column=9).value = enr.cert_number or ''
        ws_dane.cell(row=r, column=10).value = (
            enr.cert_date.strftime('%d.%m.%Y') if enr.cert_date else ''
        )

    # Arkusz1 ma {{p_lp}} – wypełnij wiersze uczestników
    _xlsx_fill_enrollment_rows(wb['Arkusz1'], enrollments, course, ctx)


def _resolve_xlsx(doc_name, instructor_count):
    """Zwraca ścieżkę do pliku: szuka {name}.{n}.xlsx, fallback do {name}.xlsx."""
    variant = TEMPLATES_DIR / f'{doc_name}.{instructor_count}.xlsx'
    if variant.exists():
        return variant
    base = TEMPLATES_DIR / f'{doc_name}.xlsx'
    if base.exists():
        return base
    return None


@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_xlsx(request, course_id, doc_name):
    if doc_name not in ALLOWED_XLSX_TEMPLATES:
        return Response({'detail': 'Nieznany dokument.'}, status=404)

    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    instructor_count = len(course.instructor_order) if course.instructor_order else course.instructors.count()
    tpl_path = _resolve_xlsx(doc_name, instructor_count)
    if tpl_path is None:
        return Response({'detail': 'Brak pliku szablonu.'}, status=404)

    wb = openpyxl.load_workbook(tpl_path)
    ctx = _build_context(course)

    if doc_name == 'zestawienie-rec':
        enrollments = list(
            course.enrollments.filter(deleted_at__isnull=True)
            .order_by('last_name', 'first_name')
        )
        for ws in wb.worksheets:
            _xlsx_fill_enrollment_rows(ws, enrollments, course, ctx)
    elif doc_name == 'obsluga-egzaminu-rec':
        enrollments = list(
            course.enrollments.filter(deleted_at__isnull=True)
            .order_by('created_at')
        )
        _xlsx_fill_obsluga_egzaminu_rec(wb, enrollments, course, ctx)
    else:
        for ws in wb.worksheets:
            _xlsx_replace(ws, ctx)

    for ws in wb.worksheets:
        ws.page_setup.orientation = 'landscape'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{doc_name}_kurs_{course_id}.xlsx"'
    return response


def _resolve_template(doc_name, instructor_count):
    """Zwraca ścieżkę do pliku: szuka {name}.{n}.docx, fallback do {name}.docx."""
    variant = TEMPLATES_DIR / f'{doc_name}.{instructor_count}.docx'
    if variant.exists():
        return variant
    base = TEMPLATES_DIR / f'{doc_name}.docx'
    if base.exists():
        return base
    return None


def _build_zaliczenia_enrollment_context(enrollment):
    """Kontekst per-uczestnik dla zaliczenia_tematow_KPP: t1_d … t15_d."""
    ctx = _build_certificate_context(enrollment)

    topics = list(Topic.objects.order_by('order', 'id'))
    passed_dict = {}
    if enrollment.user_id:
        for attempt in (
            TopicQuizAttempt.objects
            .filter(user_id=enrollment.user_id, topic__in=topics, passed=True)
            .order_by('topic_id', 'attempted_at')
        ):
            if attempt.topic_id not in passed_dict:
                passed_dict[attempt.topic_id] = attempt.attempted_at.strftime('%d.%m.%Y')

    for n, t in enumerate(topics, 1):
        ctx[f't{n}_d'] = passed_dict.get(t.id, '')
    for n in range(len(topics) + 1, 16):
        ctx[f't{n}_d'] = ''

    return ctx


@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_zaliczenia_zip(request, course_id, doc_name):
    if doc_name not in ALLOWED_ENROLLMENT_ZIP_TEMPLATES:
        return Response({'detail': 'Nieznany dokument.'}, status=404)

    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    tpl_path = TEMPLATES_DIR / f'{doc_name}.docx'
    if not tpl_path.exists():
        return Response({'detail': 'Brak pliku szablonu.'}, status=404)

    enrollments = list(
        course.enrollments.filter(deleted_at__isnull=True)
        .select_related('user', 'course')
        .order_by('last_name', 'first_name')
    )
    if not enrollments:
        return Response({'detail': 'Brak uczestników na tym kursie.'}, status=404)

    buf = BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for enrollment in enrollments:
            tpl = DocxTemplate(tpl_path)
            tpl.render(_build_zaliczenia_enrollment_context(enrollment))
            doc_buf = BytesIO()
            tpl.save(doc_buf)
            safe = f'{enrollment.last_name}_{enrollment.first_name}'.replace(' ', '_')
            zf.writestr(f'zaliczenia_{safe}.docx', doc_buf.getvalue())

    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="zaliczenia_kurs_{course_id}.zip"'
    return response


@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_document(request, course_id, doc_name):
    if doc_name not in ALLOWED_TEMPLATES:
        return Response({'detail': 'Nieznany dokument.'}, status=404)

    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    instructor_count = len(course.instructor_order) if course.instructor_order else course.instructors.count()
    tpl_path = _resolve_template(doc_name, instructor_count)
    if tpl_path is None:
        return Response({'detail': 'Brak pliku szablonu.'}, status=404)

    tpl = DocxTemplate(tpl_path)
    tpl.render(_build_context(course))

    buf = BytesIO()
    tpl.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="{doc_name}_kurs_{course_id}.docx"'
    return response


@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_document_pdf(request, course_id, doc_name):
    if doc_name not in ALLOWED_TEMPLATES:
        return Response({'detail': 'Nieznany dokument.'}, status=404)

    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    instructor_count = len(course.instructor_order) if course.instructor_order else course.instructors.count()
    tpl_path = _resolve_template(doc_name, instructor_count)
    if tpl_path is None:
        return Response({'detail': 'Brak pliku szablonu.'}, status=404)

    tpl = DocxTemplate(tpl_path)
    tpl.render(_build_context(course))

    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, f'{doc_name}.docx')
        pdf_path  = os.path.join(tmpdir, f'{doc_name}.pdf')
        tpl.save(docx_path)

        result = subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', tmpdir, docx_path],
            capture_output=True, timeout=30,
        )
        if result.returncode != 0 or not os.path.exists(pdf_path):
            return Response({'detail': 'Błąd konwersji do PDF.'}, status=500)

        with open(pdf_path, 'rb') as f:
            pdf_bytes = f.read()

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{doc_name}_kurs_{course_id}.pdf"'
    return response


def _build_certificate_context(enrollment):
    course = enrollment.course
    ctx = _build_context(course) if course else {}

    def fmt(date):
        return date.strftime('%d.%m.%Y') if date else ''

    if course:
        ids = list(
            course.enrollments.filter(is_deleted=False)
            .order_by('created_at')
            .values_list('id', flat=True)
        )
        lp = str(ids.index(enrollment.id) + 1).zfill(2) if enrollment.id in ids else ''
    else:
        lp = ''

    address_parts = [enrollment.street, enrollment.house_number]
    if enrollment.apartment_number:
        address_parts[-1] += f'/{enrollment.apartment_number}'
    full_address = f'{" ".join(address_parts)}, {enrollment.zip_code} {enrollment.city}'

    ctx.update({
        'lp': lp,
        'p_first_name':       enrollment.first_name,
        'p_last_name':        enrollment.last_name,
        'p_full_name':        f'{enrollment.first_name} {enrollment.last_name}',
        'p_pesel':            enrollment.pesel or '',
        'p_birth_date':       fmt(enrollment.birth_date),
        'p_email':            enrollment.email or '',
        'p_phone':            enrollment.phone or '',
        'p_zip_code':         enrollment.zip_code or '',
        'p_city':             enrollment.city or '',
        'p_street':           enrollment.street or '',
        'p_house_number':     enrollment.house_number or '',
        'p_apartment_number': enrollment.apartment_number or '',
        'p_address':          full_address,
        'p_cert_number':      enrollment.cert_number or '',
        'p_cert_date':        fmt(enrollment.cert_date),
    })
    return ctx


@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_certificate(request, enrollment_id, doc_name):
    if doc_name not in ALLOWED_CERT_TEMPLATES:
        return Response({'detail': 'Nieznany szablon certyfikatu.'}, status=404)

    try:
        enrollment = Enrollment.objects.select_related('course').get(pk=enrollment_id)
    except Enrollment.DoesNotExist:
        return Response({'detail': 'Uczestnik nie istnieje.'}, status=404)

    is_recert = enrollment.course and enrollment.course.course_type == 'recert'
    variant = f'{doc_name}.r.docx' if is_recert else f'{doc_name}.docx'
    tpl_path = TEMPLATES_DIR / variant
    if not tpl_path.exists():
        tpl_path = TEMPLATES_DIR / f'{doc_name}.docx'
    if not tpl_path.exists():
        return Response({'detail': 'Brak pliku szablonu certyfikatu.'}, status=404)

    tpl = DocxTemplate(tpl_path)
    tpl.render(_build_certificate_context(enrollment))

    buf = BytesIO()
    tpl.save(buf)
    buf.seek(0)

    safe_name = f'{enrollment.last_name}_{enrollment.first_name}'.replace(' ', '_')
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="{doc_name}_{safe_name}.docx"'
    return response


def _build_enrollment_context(enrollment, course_ctx, lp):
    def fmt(date):
        return date.strftime('%d.%m.%Y') if date else ''

    street = enrollment.street or ''
    house  = enrollment.house_number or ''
    if enrollment.apartment_number:
        house += f'/{enrollment.apartment_number}'
    full_address = f'{street} {house}, {enrollment.zip_code or ""} {enrollment.city or ""}'.strip(', ')

    return {
        **course_ctx,
        'p_lp':               str(lp),
        'p_first_name':       enrollment.first_name or '',
        'p_last_name':        enrollment.last_name or '',
        'p_full_name':        f'{enrollment.first_name or ""} {enrollment.last_name or ""}'.strip(),
        'p_pesel':            enrollment.pesel or '',
        'p_birth_date':       fmt(enrollment.birth_date),
        'p_email':            enrollment.email or '',
        'p_phone':            enrollment.phone or '',
        'p_zip_code':         enrollment.zip_code or '',
        'p_city':             enrollment.city or '',
        'p_street':           enrollment.street or '',
        'p_house_number':     enrollment.house_number or '',
        'p_apartment_number': enrollment.apartment_number or '',
        'p_address':          full_address,
        'p_cert_number':      enrollment.cert_number or '',
        'p_cert_date':        fmt(enrollment.cert_date) if enrollment.cert_date else '',
    }


@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_xlsx_per_enrollment(request, course_id, doc_name):
    if doc_name not in ALLOWED_ATTENDANCE_XLSX_TEMPLATES:
        return Response({'detail': 'Nieznany dokument.'}, status=404)

    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    tpl_path = TEMPLATES_DIR / f'{doc_name}.xlsx'
    if not tpl_path.exists():
        return Response({'detail': 'Brak pliku szablonu.'}, status=404)

    enrollments = list(
        course.enrollments.filter(deleted_at__isnull=True).order_by('last_name', 'first_name')
    )
    if not enrollments:
        return Response({'detail': 'Brak uczestników na tym kursie.'}, status=404)

    course_ctx = _build_context(course)
    wb = openpyxl.load_workbook(tpl_path)
    ws_tpl = wb.active

    for lp, enrollment in enumerate(enrollments, start=1):
        ws = wb.copy_worksheet(ws_tpl)
        raw_title = f'{enrollment.last_name} {enrollment.first_name}'
        safe_title = re.sub(r'[\/\?\*\[\]:\\\']', '_', raw_title)[:31]
        ws.title = safe_title
        ctx = _build_enrollment_context(enrollment, course_ctx, lp)
        _xlsx_replace(ws, ctx)

    wb.remove(ws_tpl)

    for ws in wb.worksheets:
        ws.page_setup.orientation = 'landscape'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{doc_name}_kurs_{course_id}.xlsx"'
    return response


# ─── Zbiorcze certyfikaty ZIP ─────────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAdminUser])
def download_certificates_zip(request, course_id):
    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    enrollments = list(
        course.enrollments.filter(deleted_at__isnull=True).order_by('last_name', 'first_name')
    )
    if not enrollments:
        return Response({'detail': 'Brak uczestników na tym kursie.'}, status=404)

    is_recert = course.course_type == 'recert'
    tpl_name = 'certyfikat.r.docx' if is_recert else 'certyfikat.docx'
    tpl_path = TEMPLATES_DIR / tpl_name
    if not tpl_path.exists():
        tpl_path = TEMPLATES_DIR / 'certyfikat.docx'
    if not tpl_path.exists():
        return Response({'detail': 'Brak pliku szablonu certyfikatu.'}, status=404)

    buf = BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for enrollment in enrollments:
            tpl = DocxTemplate(tpl_path)
            tpl.render(_build_certificate_context(enrollment))
            doc_buf = BytesIO()
            tpl.save(doc_buf)
            safe = f'{enrollment.last_name}_{enrollment.first_name}'.replace(' ', '_')
            zf.writestr(f'certyfikat_{safe}.docx', doc_buf.getvalue())

    buf.seek(0)
    response = HttpResponse(buf.read(), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="certyfikaty_kurs_{course_id}.zip"'
    return response


# ─── Tematy i pliki (admin) ───────────────────────────────────────────

def _serialize_question(q):
    return {
        'id':      q.id,
        'text':    q.text,
        'order':   q.order,
        'choices': [
            {'id': c.id, 'text': c.text, 'is_correct': c.is_correct, 'order': c.order}
            for c in q.choices.all()
        ],
    }


def _topic_to_dict(topic, include_questions=False):
    d = {
        'id':             topic.id,
        'title':          topic.title,
        'order':          topic.order,
        'quiz_enabled':   topic.quiz_enabled,
        'question_count': topic.questions.count() if not include_questions else None,
        'files': [
            {
                'id':          tf.id,
                'title':       tf.title,
                'uploaded_at': tf.uploaded_at.isoformat(),
                'order':       tf.order,
            }
            for tf in topic.files.all()
        ],
    }
    if include_questions:
        d['question_count'] = len(topic.questions.all())
        d['questions'] = [_serialize_question(q) for q in topic.questions.all()]
    return d


@api_view(['GET', 'POST'])
@permission_classes([IsAdminUser])
def admin_topics(request):
    if request.method == 'GET':
        topics = Topic.objects.prefetch_related('files', 'questions__choices').all()
        return Response([_topic_to_dict(t, include_questions=True) for t in topics])

    title = request.data.get('title', '').strip()
    if not title:
        return Response({'detail': 'Podaj tytuł działu.'}, status=400)
    max_order = Topic.objects.count()
    topic = Topic.objects.create(title=title, order=max_order)
    return Response(_topic_to_dict(topic), status=201)


@api_view(['PATCH', 'DELETE'])
@permission_classes([IsAdminUser])
def admin_topic_detail(request, topic_id):
    try:
        topic = Topic.objects.prefetch_related('files', 'questions__choices').get(pk=topic_id)
    except Topic.DoesNotExist:
        return Response({'detail': 'Dział nie istnieje.'}, status=404)

    if request.method == 'DELETE':
        for tf in topic.files.all():
            tf.file.delete(save=False)
        topic.delete()
        return Response(status=204)

    if 'title' in request.data:
        topic.title = request.data['title'].strip() or topic.title
    if 'order' in request.data:
        topic.order = int(request.data['order'])
    if 'quiz_enabled' in request.data:
        topic.quiz_enabled = bool(request.data['quiz_enabled'])
    topic.save()
    return Response(_topic_to_dict(topic, include_questions=True))


@api_view(['POST'])
@permission_classes([IsAdminUser])
def admin_topic_file_upload(request, topic_id):
    try:
        topic = Topic.objects.get(pk=topic_id)
    except Topic.DoesNotExist:
        return Response({'detail': 'Dział nie istnieje.'}, status=404)

    f = request.FILES.get('file')
    if not f or not f.name.lower().endswith('.pdf'):
        return Response({'detail': 'Wymagany plik PDF.'}, status=400)
    if f.content_type not in ('application/pdf', 'application/octet-stream'):
        return Response({'detail': 'Plik musi być w formacie PDF.'}, status=400)
    title = request.data.get('title', '').strip() or f.name
    max_order = topic.files.count()
    tf = TopicFile.objects.create(topic=topic, title=title, file=f, order=max_order)
    return Response({
        'id':          tf.id,
        'title':       tf.title,
        'uploaded_at': tf.uploaded_at.isoformat(),
        'order':       tf.order,
    }, status=201)


@api_view(['PATCH', 'DELETE'])
@permission_classes([IsAdminUser])
def admin_topic_file_detail(request, file_id):
    try:
        tf = TopicFile.objects.select_related('topic').get(pk=file_id)
    except TopicFile.DoesNotExist:
        return Response({'detail': 'Plik nie istnieje.'}, status=404)

    if request.method == 'DELETE':
        tf.file.delete(save=False)
        tf.delete()
        return Response(status=204)

    if 'title' in request.data:
        tf.title = request.data['title'].strip() or tf.title
    if 'order' in request.data:
        tf.order = int(request.data['order'])
    tf.save()
    return Response({'id': tf.id, 'title': tf.title, 'uploaded_at': tf.uploaded_at.isoformat(), 'order': tf.order})


# ─── Tematy i pliki (uczestnik) ───────────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def participant_topics(request):
    topics = Topic.objects.prefetch_related('files').all()
    return Response([_topic_to_dict(t) for t in topics])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def participant_progress(request):
    ids = list(TopicFileProgress.objects.filter(user=request.user).values_list('file_id', flat=True))
    return Response(ids)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def participant_toggle_progress(request, file_id):
    try:
        tf = TopicFile.objects.get(pk=file_id)
    except TopicFile.DoesNotExist:
        return Response({'detail': 'Plik nie istnieje.'}, status=404)

    obj, created = TopicFileProgress.objects.get_or_create(user=request.user, file=tf)
    if not created:
        obj.delete()
        return Response({'completed': False})
    return Response({'completed': True}, status=201)


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def quiz_progress(request):
    if request.method == 'GET':
        qs = QuizProgress.objects.filter(user=request.user)
        return Response({obj.category_id: obj.last_index for obj in qs})

    cat_id     = request.data.get('category_id', '').strip()
    last_index = request.data.get('last_index')
    if not cat_id or last_index is None:
        return Response({'detail': 'category_id i last_index są wymagane.'}, status=400)
    if not isinstance(last_index, int) or last_index < 0:
        return Response({'detail': 'last_index musi być liczbą całkowitą >= 0.'}, status=400)

    obj, _ = QuizProgress.objects.update_or_create(
        user=request.user,
        category_id=cat_id,
        defaults={'last_index': last_index},
    )
    return Response({'category_id': obj.category_id, 'last_index': obj.last_index})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def participant_topic_file(request, file_id):
    try:
        tf = TopicFile.objects.get(pk=file_id)
    except TopicFile.DoesNotExist:
        return Response({'detail': 'Plik nie istnieje.'}, status=404)

    try:
        f = tf.file.open('rb')
    except (FileNotFoundError, OSError):
        return Response({'detail': 'Plik nie istnieje na serwerze.'}, status=404)

    safe_title = tf.title.replace(' ', '_')
    response = FileResponse(f, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{safe_title}.pdf"'
    return response


# ─── Pliki kursu (upload przez admina) ───────────────────────────────

@api_view(['GET', 'POST'])
@permission_classes([IsAdminUser])
def course_file_list(request, course_id):
    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        return Response({'detail': 'Kurs nie istnieje.'}, status=404)

    if request.method == 'GET':
        files = course.uploaded_files.all()
        return Response([_course_file_dict(f) for f in files])

    f = request.FILES.get('file')
    if not f:
        return Response({'detail': 'Wymagany plik.'}, status=400)
    label = request.data.get('label', '').strip() or f.name
    cf = CourseFile.objects.create(course=course, label=label, file=f)
    return Response(_course_file_dict(cf), status=201)


@api_view(['GET', 'DELETE'])
@permission_classes([IsAdminUser])
def course_file_detail(request, file_id):
    try:
        cf = CourseFile.objects.get(pk=file_id)
    except CourseFile.DoesNotExist:
        return Response({'detail': 'Plik nie istnieje.'}, status=404)

    if request.method == 'DELETE':
        cf.file.delete(save=False)
        cf.delete()
        return Response(status=204)

    try:
        f = cf.file.open('rb')
    except (FileNotFoundError, OSError):
        return Response({'detail': 'Plik nie istnieje na serwerze.'}, status=404)

    import mimetypes
    mime, _ = mimetypes.guess_type(cf.file.name)
    mime = mime or 'application/octet-stream'
    safe_label = cf.label.replace(' ', '_')
    ext = os.path.splitext(cf.file.name)[1]
    response = FileResponse(f, content_type=mime)
    response['Content-Disposition'] = f'attachment; filename="{safe_label}{ext}"'
    return response


def _course_file_dict(cf):
    return {
        'id':          cf.id,
        'label':       cf.label,
        'filename':    os.path.basename(cf.file.name),
        'uploaded_at': cf.uploaded_at.isoformat(),
    }


# ─── Pytania zaliczeniowe (admin) ────────────────────────────────────

@api_view(['GET', 'POST'])
@permission_classes([IsAdminUser])
def admin_topic_questions(request, topic_id):
    try:
        topic = Topic.objects.get(pk=topic_id)
    except Topic.DoesNotExist:
        return Response({'detail': 'Dział nie istnieje.'}, status=404)

    if request.method == 'GET':
        questions = topic.questions.prefetch_related('choices').all()
        return Response([_serialize_question(q) for q in questions])

    if topic.questions.count() >= 5:
        return Response({'detail': 'Maksymalna liczba pytań (5) osiągnięta.'}, status=400)

    text = request.data.get('text', '').strip()
    if not text:
        return Response({'detail': 'Treść pytania jest wymagana.'}, status=400)
    choices_data = request.data.get('choices', [])
    if len(choices_data) < 2:
        return Response({'detail': 'Wymagane co najmniej 2 odpowiedzi.'}, status=400)
    if not any(c.get('is_correct') for c in choices_data):
        return Response({'detail': 'Zaznacz poprawną odpowiedź.'}, status=400)

    order = topic.questions.count()
    q = TopicQuestion.objects.create(topic=topic, text=text, order=order)
    for i, c in enumerate(choices_data):
        TopicAnswerChoice.objects.create(question=q, text=c.get('text', '').strip(), is_correct=bool(c.get('is_correct')), order=i)
    return Response(_serialize_question(q), status=201)


@api_view(['PATCH', 'DELETE'])
@permission_classes([IsAdminUser])
def admin_topic_question_detail(request, question_id):
    try:
        q = TopicQuestion.objects.prefetch_related('choices').get(pk=question_id)
    except TopicQuestion.DoesNotExist:
        return Response({'detail': 'Pytanie nie istnieje.'}, status=404)

    if request.method == 'DELETE':
        q.delete()
        return Response(status=204)

    if 'text' in request.data:
        q.text = request.data['text'].strip() or q.text
        q.save()
    if 'choices' in request.data:
        choices_data = request.data['choices']
        if not any(c.get('is_correct') for c in choices_data):
            return Response({'detail': 'Zaznacz poprawną odpowiedź.'}, status=400)
        q.choices.all().delete()
        for i, c in enumerate(choices_data):
            TopicAnswerChoice.objects.create(question=q, text=c.get('text', '').strip(), is_correct=bool(c.get('is_correct')), order=i)
    return Response(_serialize_question(q))


# ─── Pytania zaliczeniowe (uczestnik) ────────────────────────────────

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def participant_topic_quiz(request, topic_id):
    try:
        topic = Topic.objects.prefetch_related('questions__choices').get(pk=topic_id)
    except Topic.DoesNotExist:
        return Response({'detail': 'Dział nie istnieje.'}, status=404)

    if not topic.quiz_enabled:
        return Response({'detail': 'Quiz nie jest dostępny.'}, status=403)

    data = []
    has_passed = TopicQuizAttempt.objects.filter(user=request.user, topic=topic, passed=True).exists()
    for q in topic.questions.all():
        data.append({
            'id':      q.id,
            'text':    q.text,
            'choices': [
                {'id': c.id, 'text': c.text, **(({'is_correct': c.is_correct}) if has_passed else {})}
                for c in q.choices.all()
            ],
        })
    return Response(data)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def participant_submit_quiz(request, topic_id):
    try:
        topic = Topic.objects.prefetch_related('questions__choices').get(pk=topic_id)
    except Topic.DoesNotExist:
        return Response({'detail': 'Dział nie istnieje.'}, status=404)

    if not topic.quiz_enabled:
        return Response({'detail': 'Quiz nie jest dostępny.'}, status=403)

    questions = list(topic.questions.all())
    if not questions:
        return Response({'detail': 'Brak pytań w tym dziale.'}, status=400)

    answers = request.data.get('answers', {})
    score = 0
    results = []
    for q in questions:
        chosen_id = answers.get(str(q.id))
        correct = q.choices.filter(is_correct=True).first()
        is_correct = False
        if chosen_id:
            try:
                chosen = q.choices.get(pk=int(chosen_id))
                is_correct = chosen.is_correct
            except (TopicAnswerChoice.DoesNotExist, ValueError):
                pass
        if is_correct:
            score += 1
        results.append({
            'question_id':      q.id,
            'correct':          is_correct,
            'correct_choice_id': correct.id if correct else None,
        })

    total = len(questions)
    passed = score == total
    attempt = TopicQuizAttempt.objects.create(user=request.user, topic=topic, score=score, total=total, passed=passed)
    return Response({'score': score, 'total': total, 'passed': passed, 'results': results, 'attempted_at': attempt.attempted_at.isoformat()})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def participant_quiz_results(request):
    attempts = TopicQuizAttempt.objects.filter(user=request.user).order_by('topic_id', '-passed', '-score', '-attempted_at')
    best = {}
    for a in attempts:
        if a.topic_id not in best:
            best[a.topic_id] = {
                'topic_id':    a.topic_id,
                'score':       a.score,
                'total':       a.total,
                'passed':      a.passed,
                'attempted_at': a.attempted_at.isoformat(),
            }
    return Response(list(best.values()))
